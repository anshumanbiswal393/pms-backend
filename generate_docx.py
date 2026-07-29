import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_workflow_docx():
    doc = Document()

    # Set Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles & Fonts
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Slate 700

    # Helper function to style headings
    def add_custom_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E) # Teal 700
        p.paragraph_format.space_after = Pt(4)

    def add_custom_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B) # Slate 500
        p.paragraph_format.space_after = Pt(20)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Dark Blue

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E) # Teal 700

    def add_callout(title, body_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)
        
        # Set shading background (Light Teal F0FDF4)
        shading = parse_xml(r'<w:shd {} w:fill="F0FDF4"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        
        # Set left border thick Teal 700
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(r'''
            <w:tcBorders {} >
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="0F766E"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        '''.format(nsdecls('w')))
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        
        run_title = p.add_run(f"📌 {title}\n")
        run_title.font.bold = True
        run_title.font.size = Pt(11)
        run_title.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)

        run_body = p.add_run(body_text)
        run_body.font.size = Pt(10.5)
        run_body.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Document Header
    add_custom_title("Retrod PMS — End-to-End System Operational & Testing Workflow Guide")
    add_custom_subtitle("Step-by-Step Functional Testing, SuperAdmin Setup, Tenant Onboarding, Inventory Layout, Commercial Services, Reservation Lifecycle, and Housekeeping Turnaround")

    add_callout(
        "Workflow Document Scope",
        "This guide details the complete sequential testing flow for Retrod PMS — starting from initial SuperAdmin platform configuration to tenant onboarding, room & physical layout setup, commercial catalog, guest reservations, check-in, folio billing, check-out, and housekeeping room turnaround."
    )

    # Section 1: Overview Matrix Table
    add_heading_1("1. End-to-End System Execution Sequence Overview")

    phases_data = [
        ("Phase 1: SuperAdmin Setup", "Platform Config, Currencies, Taxes, Docs & Languages", "http://localhost:5173/superadmin/settings"),
        ("Phase 2: Partner Onboarding", "Create Tenant Partner, Owner Account & Register Property", "http://localhost:5173/superadmin/tenants/new"),
        ("Phase 3: Tenant Setup", "Tenant Owner Login, Select Property Context & Set Policies", "http://localhost:5173/settings"),
        ("Phase 4: Rooms & Inventory", "Create Room Types, Buildings, Floors & Room Assignments", "http://localhost:5173/rooms"),
        ("Phase 5: Commercial Catalog", "Setup Service Categories, Add-On Services & Rate Plans", "http://localhost:5173/services"),
        ("Phase 6: Guest Lifecycle", "Guest Profile, Booking, Check-In, Folio Billing & Check-Out", "http://localhost:5173/reservations"),
        ("Phase 7: Housekeeping", "Auto Dirty Mark, Cleaning Task Assignment & Room Release", "http://localhost:5173/housekeeping")
    ]

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_titles = ["Phase / Module", "Primary Functional Scope", "UI Screen URL"]
    for i, title in enumerate(hdr_titles):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shading = parse_xml(r'<w:shd {} w:fill="0F766E"/>'.format(nsdecls('w')))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)

    for phase, scope, url in phases_data:
        row_cells = table.add_row().cells
        row_cells[0].text = phase
        row_cells[1].text = scope
        row_cells[2].text = url
        row_cells[0].paragraphs[0].runs[0].font.bold = True

    # Adjust widths
    col_widths = [Inches(2.0), Inches(3.0), Inches(2.0)]
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Phase 1
    add_heading_1("2. Phase 1: SuperAdmin Platform Setup")
    
    add_heading_2("Step 1.1: SuperAdmin Login")
    p = doc.add_paragraph()
    p.add_run("• Login URL: ").bold = True
    p.add_run("http://localhost:5173/login\n")
    p.add_run("• SuperAdmin Credentials: ").bold = True
    p.add_run("admin@retrod.in / AdminPassword123!\n")
    p.add_run("• Operational Verification: ").bold = True
    p.add_run("SuperAdmin authenticates platform-wide with full permissions (['*:*']).")

    add_heading_2("Step 1.2: Global System Configuration")
    p = doc.add_paragraph()
    p.add_run("1. Currency Settings: ").bold = True
    p.add_run("Click '+ Add Currency' slide-over right-side sheet modal. Select Country of Origin (e.g. United States, India, United Arab Emirates) -> Save currency (USD, INR, EUR, AED). Set default application currency.\n")
    p.add_run("2. Global Tax Setup: ").bold = True
    p.add_run("Configure tax rules (VAT 15%, Service Tax 5%, Tourism Levy 10.00 Fixed).\n")
    p.add_run("3. Document Identification Setup: ").bold = True
    p.add_run("Define accepted guest identification types (Passport, National ID / Aadhaar, Driver's License) and check-in rules.\n")
    p.add_run("4. Languages & Formats: ").bold = True
    p.add_run("Set default date formats (YYYY-MM-DD), time formats (HH:mm), and active system languages.")

    # Phase 2
    add_heading_1("3. Phase 2: Partner & Property Onboarding")
    
    add_heading_2("Step 2.1: Onboard Hospitality Partner (Tenant)")
    p = doc.add_paragraph()
    p.add_run("• Navigation: ").bold = True
    p.add_run("http://localhost:5173/superadmin/tenants/new\n")
    p.add_run("• Partner Details: ").bold = True
    p.add_run("Enter Partner Name (e.g., 'Rajdhani Hospitality Group'), Subdomain (e.g., 'rajdhani'), and select Subscription Plan ('Standard Enterprise Plan').\n")
    p.add_run("• Owner Account: ").bold = True
    p.add_run("Provision Tenant Owner account credentials (e.g., owner@rajdhani.com / OwnerPassword123!).")

    add_heading_2("Step 2.2: Register Property under Tenant")
    p = doc.add_paragraph()
    p.add_run("• Navigation: ").bold = True
    p.add_run("http://localhost:5173/superadmin/properties/new\n")
    p.add_run("• Property Details: ").bold = True
    p.add_run("Select Tenant ('Rajdhani Hospitality Group'), enter Property Name ('Hotel Rajdhani Grand'), select Property Type ('Hotel'), address, country, and contact phone.\n")
    p.add_run("• RBAC Automatic Provisioning: ").bold = True
    p.add_run("System automatically maps tenant-wide operational RBAC permissions for the Tenant Owner across the onboarded property.")

    # Phase 3
    add_heading_1("4. Phase 3: Tenant Property Operations")
    
    add_heading_2("Step 3.1: Tenant Owner Login")
    p = doc.add_paragraph()
    p.add_run("• Credentials: ").bold = True
    p.add_run("Log in with Tenant Owner account (owner@rajdhani.com).\n")
    p.add_run("• Property Context: ").bold = True
    p.add_run("System automatically loads 'Hotel Rajdhani Grand' as active operating property.")

    add_heading_2("Step 3.2: Property Settings & Operational Policies")
    p = doc.add_paragraph()
    p.add_run("• Navigation: ").bold = True
    p.add_run("http://localhost:5173/settings\n")
    p.add_run("• Policy Setup: ").bold = True
    p.add_run("Configure check-in timing (14:00), check-out timing (11:00), cancellation terms, and active payment gateways (Cash, Card, Razorpay).")

    # Phase 4
    add_heading_1("5. Phase 4: Rooms, Inventory & Physical Layout Setup")
    
    add_heading_2("Step 4.1: Create Room Types")
    p = doc.add_paragraph()
    p.add_run("• Navigation: ").bold = True
    p.add_run("http://localhost:5173/rooms -> '+ Add Room Type'\n")
    p.add_run("• Room Types Configured:\n").bold = True
    p.add_run("  - Executive Suite: Base Rate ₹5,000 | Base Occupancy 2 | Max Occupancy 3 | Amenities: WiFi, AC, TV, Mini Bar\n")
    p.add_run("  - Deluxe King: Base Rate ₹3,500 | Base Occupancy 2 | Max Occupancy 2 | Amenities: WiFi, AC, TV\n")
    p.add_run("  - Presidential Suite: Base Rate ₹12,000 | Base Occupancy 4 | Max Occupancy 6 | Amenities: Full Luxury Kit")

    add_heading_2("Step 4.2: Building & Floor Layout Management")
    p = doc.add_paragraph()
    p.add_run("• Buildings / Wings: ").bold = True
    p.add_run("Create 'Main Wing' and 'Tower A'.\n")
    p.add_run("• Floor Setup: ").bold = True
    p.add_run("Create '1st Floor', '2nd Floor', and '3rd Floor' linked to respective buildings.")

    add_heading_2("Step 4.3: Physical Room Inventory Assignment")
    p = doc.add_paragraph()
    p.add_run("• Navigation: ").bold = True
    p.add_run("http://localhost:5173/rooms -> '+ Add Room'\n")
    p.add_run("• Room Assignment: ").bold = True
    p.add_run("Assign Room 101, 102, 103, 201, 202 linked to Room Type, Building, Floor, and set initial operational status to 'Clean / Available'.")

    # Phase 5
    add_heading_1("6. Phase 5: Commercial Catalog, Services & Packages")
    
    add_heading_2("Step 5.1: Create Service Categories & Items")
    p = doc.add_paragraph()
    p.add_run("• Navigation: ").bold = True
    p.add_run("http://localhost:5173/services -> '+ Add Service'\n")
    p.add_run("• Catalog Items Created:\n").bold = True
    p.add_run("  - Buffet Breakfast: ₹500 (Category: Dining & Room Service)\n")
    p.add_run("  - Airport Pick & Drop: ₹1,500 (Category: Transportation)\n")
    p.add_run("  - Full Body Spa Massage: ₹2,500 (Category: Wellness & Spa)")

    add_heading_2("Step 5.2: Configure Rate Plans & Packages")
    p = doc.add_paragraph()
    p.add_run("• Rate Plans: ").bold = True
    p.add_run("Configure 'Standard Flexible Rate' and 'Non-Refundable AP Rate'. Link meal plan inclusions (CP - Bed & Breakfast, MAP - Half Board, AP - Full Board).")

    # Phase 6
    add_heading_1("7. Phase 6: Guest Lifecycle & Reservation Management")
    
    add_heading_2("Step 6.1: Create Guest Profile & Reservation")
    p = doc.add_paragraph()
    p.add_run("• Navigation: ").bold = True
    p.add_run("http://localhost:5173/reservations -> '+ New Reservation'\n")
    p.add_run("• Booking Flow: ").bold = True
    p.add_run("Select stay dates -> Select Room Type ('Executive Suite') -> Assign Room 101 -> Enter Guest details ('Rajesh Kumar', +91 9876543210, rajesh@example.com) -> Select Rate Plan & Add-On Services -> Confirm Booking. Booking status transitions to 'Confirmed'.")

    add_heading_2("Step 6.2: Front Office Check-In")
    p = doc.add_paragraph()
    p.add_run("• Actions: ").bold = True
    p.add_run("Locate booking -> Click 'Check In' -> Verify Guest ID Document (Passport / Aadhaar) -> Record advance payment deposit -> Confirm Check-In. Reservation status transitions to 'Checked In', Room 101 status updates to 'Occupied'.")

    add_heading_2("Step 6.3: In-House Guest Charges & Folio Billing")
    p = doc.add_paragraph()
    p.add_run("• Folio Postings: ").bold = True
    p.add_run("Post additional room service or laundry charges directly to Room 101 Guest Folio.")

    add_heading_2("Step 6.4: Front Office Check-Out & Settlement")
    p = doc.add_paragraph()
    p.add_run("• Actions: ").bold = True
    p.add_run("Click 'Check Out' on Room 101 -> Generate Final Folio Invoice (Room charges + Services + Taxes) -> Collect payment settlement (Cash / Card / Razorpay) -> Confirm Check-Out. Reservation status transitions to 'Checked Out'. Room 101 status automatically updates to 'Dirty / Housekeeping Pending'.")

    add_heading_2("Step 6.5: Alternate Flow — Reservation Cancellation")
    p = doc.add_paragraph()
    p.add_run("• Actions: ").bold = True
    p.add_run("Select Confirmed booking -> Click 'Cancel Booking' -> System releases assigned room inventory back to availability pool.")

    # Phase 7
    add_heading_1("8. Phase 7: Housekeeping & Room Turnaround")
    p = doc.add_paragraph()
    p.add_run("• Navigation: ").bold = True
    p.add_run("http://localhost:5173/housekeeping\n")
    p.add_run("• Turnaround Flow:\n").bold = True
    p.add_run("  1. Room 101 appears automatically under 'Dirty Rooms' list post Check-Out.\n")
    p.add_run("  2. Assign cleaning task to Housekeeping staff.\n")
    p.add_run("  3. Staff cleans room and clicks 'Mark Clean'.\n")
    p.add_run("  4. Supervisor inspects room and clicks 'Mark Inspected & Available'.\n")
    p.add_run("  5. Room 101 status updates to 'Clean / Ready for next Check-in'.")

    # Save Document
    file_name = "Retrod_PMS_End_to_End_Testing_Workflow_Guide.docx"
    doc.save(file_name)
    print(f"DOCX REPORT GENERATED SUCCESSFULLY: {file_name}")

if __name__ == "__main__":
    create_workflow_docx()
